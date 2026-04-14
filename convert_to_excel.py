import json
import sys
import os
from openpyxl import Workbook
from openpyxl.styles import Font

def convert():
    # Get the full file path from command line argument
    filename = sys.argv[1] if len(sys.argv) > 1 else "reports/report.xlsx"
    
    # Ensure the directory exists
    os.makedirs(os.path.dirname(filename), exist_ok=True)
    
    if not os.path.exists('temp_data.json'):
        print("Error: temp_data.json not found")
        return

    with open('temp_data.json', 'r') as f:
        data = json.load(f)

    wb = Workbook()
    
    # --- SHEET 1: EXPENSES ---
    ws1 = wb.active
    ws1.title = "Expenses"
    if data.get('expenses'):
        # Add Headers
        headers = list(data['expenses'][0].keys())
        ws1.append(headers)
        # Add Data
        for row in data['expenses']:
            ws1.append(list(row.values()))
        # Add Total Row
        total_row = ["TOTAL", "", "", "", data['totals']['total_expenses']]
        ws1.append(total_row)
        ws1[f'A{ws1.max_row}'].font = Font(bold=True)

    # --- SHEET 2: PAYMENTS ---
    ws2 = wb.create_sheet("Payments")
    if data.get('payments'):
        # Add Headers
        headers = list(data['payments'][0].keys())
        ws2.append(headers)
        # Add Data
        for row in data['payments']:
            ws2.append(list(row.values()))
        # Add Total Row
        total_row = ["TOTAL", "", "", "", data['totals']['total_payments']]
        ws2.append(total_row)
        ws2[f'A{ws2.max_row}'].font = Font(bold=True)

    # Save to the specified path
    wb.save(filename)
    print(f"Report saved to: {filename}")
    
    # Cleanup
    if os.path.exists('temp_data.json'):
        os.remove('temp_data.json')

if __name__ == "__main__":
    convert()